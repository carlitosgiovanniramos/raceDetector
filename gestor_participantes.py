"""
Gestor de Participantes para Sistema de Cronometraje de Carreras
Maneja el almacenamiento y consulta de datos de participantes en Excel
"""

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import os
from datetime import datetime
import csv


class GestorParticipantes:
    """Gestiona la base de datos de participantes en formato Excel"""
    
    def __init__(self, archivo_excel="participantes.xlsx"):
        self.archivo = archivo_excel
        self.columnas = [
            "Cédula", "Nombre", "Apellido", "Dirección", 
            "Contacto", "Estado Pago", "Número Dorsal"
        ]
        self._inicializar_archivo()
    
    def _formatear_dorsal(self, numero_dorsal):
        """
        Formatea el número de dorsal a 3 dígitos con ceros a la izquierda
        Ejemplos: 8 -> 008, 25 -> 025, 100 -> 100
        
        Args:
            numero_dorsal: Número de dorsal (puede ser string o int)
            
        Returns:
            str: Número formateado con 3 dígitos o el valor original si está vacío
        """
        if not numero_dorsal:
            return ""
        
        # Convertir a string y quitar espacios
        dorsal_str = str(numero_dorsal).strip()
        
        # Si está vacío después de quitar espacios, retornar vacío
        if not dorsal_str:
            return ""
        
        # Intentar convertir a número y formatear
        try:
            # Quitar ceros a la izquierda primero (si los tiene)
            numero = int(dorsal_str)
            # Formatear con 3 dígitos, agregando ceros a la izquierda si es necesario
            return f"{numero:03d}"
        except ValueError:
            # Si no se puede convertir a número, retornar el valor original
            return dorsal_str
    
    def _inicializar_archivo(self):
        """Crea el archivo Excel si no existe"""
        if not os.path.exists(self.archivo):
            wb = Workbook()
            ws = wb.active
            ws.title = "Participantes"
            
            # Estilo para encabezados
            header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF", size=12)
            
            # Escribir encabezados
            for col_num, columna in enumerate(self.columnas, 1):
                cell = ws.cell(row=1, column=col_num)
                cell.value = columna
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            
            # Ajustar anchos de columna
            anchos = [15, 20, 20, 30, 15, 15, 15]
            for col_num, ancho in enumerate(anchos, 1):
                ws.column_dimensions[openpyxl.utils.get_column_letter(col_num)].width = ancho
            
            wb.save(self.archivo)
            print(f"✅ Archivo '{self.archivo}' creado exitosamente")
    
    def agregar_participante(self, cedula, nombre, apellido, direccion, contacto, 
                            estado_pago="PENDIENTE", numero_dorsal=""):
        """
        Agrega un nuevo participante al archivo Excel
        
        Args:
            cedula: Número de cédula/DNI del participante
            nombre: Nombre del participante
            apellido: Apellido del participante
            direccion: Dirección del participante
            contacto: Teléfono/email de contacto
            estado_pago: PAGADO o PENDIENTE
            numero_dorsal: Número de dorsal asignado
            
        Returns:
            tuple: (éxito: bool, mensaje: str)
        """
        # Validar que no exista la cédula
        if self.buscar_por_cedula(cedula):
            return False, f"❌ Ya existe un participante con cédula {cedula}"
        
        # Formatear el número de dorsal a 3 dígitos
        numero_dorsal = self._formatear_dorsal(numero_dorsal)
        
        # Validar que el dorsal no esté repetido (si se proporciona)
        if numero_dorsal and self.buscar_por_dorsal(numero_dorsal):
            return False, f"❌ El número de dorsal {numero_dorsal} ya está asignado"
        
        # Validar estado de pago
        estado_pago = estado_pago.upper()
        if estado_pago not in ["PAGADO", "PENDIENTE"]:
            estado_pago = "PENDIENTE"
        
        try:
            wb = openpyxl.load_workbook(self.archivo)
            ws = wb.active
            
            # Encontrar la siguiente fila vacía
            siguiente_fila = ws.max_row + 1
            
            # Escribir datos
            ws.cell(row=siguiente_fila, column=1).value = str(cedula)
            ws.cell(row=siguiente_fila, column=2).value = nombre
            ws.cell(row=siguiente_fila, column=3).value = apellido
            ws.cell(row=siguiente_fila, column=4).value = direccion
            ws.cell(row=siguiente_fila, column=5).value = contacto
            ws.cell(row=siguiente_fila, column=6).value = estado_pago
            ws.cell(row=siguiente_fila, column=7).value = str(numero_dorsal) if numero_dorsal else ""
            
            # Aplicar estilo a la fila de estado de pago
            cell_pago = ws.cell(row=siguiente_fila, column=6)
            if estado_pago == "PAGADO":
                cell_pago.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                cell_pago.font = Font(color="006100", bold=True)
            else:
                cell_pago.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                cell_pago.font = Font(color="9C0006", bold=True)
            
            # Intentar guardar el archivo
            try:
                wb.save(self.archivo)
                wb.close()
            except PermissionError:
                wb.close()
                return False, f"❌ No se puede guardar: el archivo '{self.archivo}' está abierto en otra aplicación. Ciérralo e intenta de nuevo."
            
            return True, f"✅ Participante {nombre} {apellido} agregado exitosamente"
        
        except PermissionError:
            return False, f"❌ No se puede acceder al archivo '{self.archivo}': está abierto en otra aplicación. Ciérralo e intenta de nuevo."
        except Exception as e:
            return False, f"❌ Error al agregar participante: {str(e)}"
    
    def buscar_por_cedula(self, cedula):
        """
        Busca un participante por su cédula
        
        Returns:
            dict o None: Datos del participante si existe
        """
        try:
            wb = openpyxl.load_workbook(self.archivo)
            ws = wb.active
            
            for fila in range(2, ws.max_row + 1):
                if str(ws.cell(row=fila, column=1).value) == str(cedula):
                    # Leer el dorsal y formatearlo automáticamente
                    dorsal_raw = ws.cell(row=fila, column=7).value
                    dorsal_formateado = self._formatear_dorsal(dorsal_raw) if dorsal_raw else ""
                    
                    return {
                        "fila": fila,
                        "cedula": ws.cell(row=fila, column=1).value,
                        "nombre": ws.cell(row=fila, column=2).value,
                        "apellido": ws.cell(row=fila, column=3).value,
                        "direccion": ws.cell(row=fila, column=4).value,
                        "contacto": ws.cell(row=fila, column=5).value,
                        "estado_pago": ws.cell(row=fila, column=6).value,
                        "numero_dorsal": dorsal_formateado
                    }
            return None
        except Exception as e:
            print(f"❌ Error al buscar por cédula: {str(e)}")
            return None
    
    def buscar_por_dorsal(self, numero_dorsal):
        """
        Busca un participante por su número de dorsal
        
        Returns:
            dict o None: Datos del participante si existe
        """
        try:
            # Formatear el dorsal de búsqueda para comparar
            dorsal_busqueda = self._formatear_dorsal(numero_dorsal)
            
            wb = openpyxl.load_workbook(self.archivo)
            ws = wb.active
            
            for fila in range(2, ws.max_row + 1):
                dorsal_cell = ws.cell(row=fila, column=7).value
                # Formatear el dorsal leído del Excel
                dorsal_formateado = self._formatear_dorsal(dorsal_cell) if dorsal_cell else ""
                
                if dorsal_formateado == dorsal_busqueda:
                    return {
                        "fila": fila,
                        "cedula": ws.cell(row=fila, column=1).value,
                        "nombre": ws.cell(row=fila, column=2).value,
                        "apellido": ws.cell(row=fila, column=3).value,
                        "direccion": ws.cell(row=fila, column=4).value,
                        "contacto": ws.cell(row=fila, column=5).value,
                        "estado_pago": ws.cell(row=fila, column=6).value,
                        "numero_dorsal": dorsal_formateado
                    }
            return None
        except Exception as e:
            print(f"❌ Error al buscar por dorsal: {str(e)}")
            return None
    
    def obtener_todos(self):
        """
        Obtiene todos los participantes
        
        Returns:
            list: Lista de diccionarios con datos de participantes
        """
        participantes = []
        try:
            wb = openpyxl.load_workbook(self.archivo)
            ws = wb.active
            
            for fila in range(2, ws.max_row + 1):
                # Leer el dorsal y formatearlo automáticamente
                dorsal_raw = ws.cell(row=fila, column=7).value
                dorsal_formateado = self._formatear_dorsal(dorsal_raw) if dorsal_raw else ""
                
                participante = {
                    "fila": fila,
                    "cedula": ws.cell(row=fila, column=1).value,
                    "nombre": ws.cell(row=fila, column=2).value,
                    "apellido": ws.cell(row=fila, column=3).value,
                    "direccion": ws.cell(row=fila, column=4).value,
                    "contacto": ws.cell(row=fila, column=5).value,
                    "estado_pago": ws.cell(row=fila, column=6).value,
                    "numero_dorsal": dorsal_formateado
                }
                participantes.append(participante)
            
            return participantes
        except Exception as e:
            print(f"❌ Error al obtener participantes: {str(e)}")
            return []
    
    def actualizar_participante(self, cedula, **campos):
        """
        Actualiza los datos de un participante existente
        
        Args:
            cedula: Cédula del participante a actualizar
            **campos: Campos a actualizar (nombre, apellido, direccion, contacto, estado_pago, numero_dorsal)
        
        Returns:
            tuple: (éxito: bool, mensaje: str)
        """
        participante = self.buscar_por_cedula(cedula)
        if not participante:
            return False, f"❌ No existe un participante con cédula {cedula}"
        
        try:
            wb = openpyxl.load_workbook(self.archivo)
            ws = wb.active
            fila = participante["fila"]
            
            # Mapeo de campos a columnas
            mapeo = {
                "nombre": 2,
                "apellido": 3,
                "direccion": 4,
                "contacto": 5,
                "estado_pago": 6,
                "numero_dorsal": 7
            }
            
            # Formatear el dorsal si se está actualizando
            if "numero_dorsal" in campos:
                campos["numero_dorsal"] = self._formatear_dorsal(campos["numero_dorsal"])
            
            # Validar si se está cambiando el dorsal y ya existe
            if "numero_dorsal" in campos and campos["numero_dorsal"]:
                nuevo_dorsal = str(campos["numero_dorsal"])
                dorsal_existente = self.buscar_por_dorsal(nuevo_dorsal)
                if dorsal_existente and str(dorsal_existente["cedula"]) != str(cedula):
                    return False, f"❌ El dorsal {nuevo_dorsal} ya está asignado a otro participante"
            
            # Actualizar campos
            for campo, valor in campos.items():
                if campo in mapeo:
                    col = mapeo[campo]
                    if campo == "estado_pago":
                        valor = valor.upper()
                        cell = ws.cell(row=fila, column=col)
                        cell.value = valor
                        # Aplicar estilo
                        if valor == "PAGADO":
                            cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                            cell.font = Font(color="006100", bold=True)
                        else:
                            cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                            cell.font = Font(color="9C0006", bold=True)
                    else:
                        ws.cell(row=fila, column=col).value = str(valor) if valor else ""
            
            # Intentar guardar el archivo
            try:
                wb.save(self.archivo)
                wb.close()
            except PermissionError:
                wb.close()
                return False, f"❌ No se puede guardar: el archivo '{self.archivo}' está abierto en otra aplicación. Ciérralo e intenta de nuevo."
            
            return True, f"✅ Participante actualizado exitosamente"
        
        except PermissionError:
            return False, f"❌ No se puede acceder al archivo '{self.archivo}': está abierto en otra aplicación. Ciérralo e intenta de nuevo."
        except Exception as e:
            return False, f"❌ Error al actualizar participante: {str(e)}"
    
    def eliminar_participante(self, cedula):
        """
        Elimina un participante del registro
        
        Returns:
            tuple: (éxito: bool, mensaje: str)
        """
        participante = self.buscar_por_cedula(cedula)
        if not participante:
            return False, f"❌ No existe un participante con cédula {cedula}"
        
        try:
            wb = openpyxl.load_workbook(self.archivo)
            ws = wb.active
            
            ws.delete_rows(participante["fila"])
            
            wb.save(self.archivo)
            return True, f"✅ Participante eliminado exitosamente"
        
        except Exception as e:
            return False, f"❌ Error al eliminar participante: {str(e)}"
    
    def importar_csv(self, archivo_csv):
        """
        Importa participantes desde un archivo CSV
        
        Returns:
            tuple: (éxito: bool, mensaje: str, importados: int, errores: list)
        """
        if not os.path.exists(archivo_csv):
            return False, f"❌ No se encontró el archivo {archivo_csv}", 0, []
        
        importados = 0
        errores = []
        
        try:
            with open(archivo_csv, 'r', encoding='utf-8') as f:
                reader = csv.DictReader(f)
                
                for idx, row in enumerate(reader, start=2):
                    try:
                        cedula = row.get('cedula', '').strip()
                        nombre = row.get('nombre', '').strip()
                        apellido = row.get('apellido', '').strip()
                        direccion = row.get('direccion', '').strip()
                        contacto = row.get('contacto', '').strip()
                        estado_pago = row.get('estado_pago', 'PENDIENTE').strip().upper()
                        numero_dorsal = row.get('numero_dorsal', '').strip()
                        
                        if not cedula or not nombre:
                            errores.append(f"Fila {idx}: Falta cédula o nombre")
                            continue
                        
                        exito, msg = self.agregar_participante(
                            cedula, nombre, apellido, direccion, 
                            contacto, estado_pago, numero_dorsal
                        )
                        
                        if exito:
                            importados += 1
                        else:
                            errores.append(f"Fila {idx}: {msg}")
                    
                    except Exception as e:
                        errores.append(f"Fila {idx}: {str(e)}")
            
            mensaje = f"✅ Importación completada: {importados} participantes importados"
            if errores:
                mensaje += f", {len(errores)} errores"
            
            return True, mensaje, importados, errores
        
        except Exception as e:
            return False, f"❌ Error al leer CSV: {str(e)}", 0, []
    
    def exportar_csv(self, archivo_salida="participantes_export.csv"):
        """
        Exporta todos los participantes a un archivo CSV
        
        Returns:
            tuple: (éxito: bool, mensaje: str)
        """
        try:
            participantes = self.obtener_todos()
            
            with open(archivo_salida, 'w', newline='', encoding='utf-8') as f:
                if participantes:
                    fieldnames = ['cedula', 'nombre', 'apellido', 'direccion', 
                                'contacto', 'estado_pago', 'numero_dorsal']
                    writer = csv.DictWriter(f, fieldnames=fieldnames)
                    
                    writer.writeheader()
                    for p in participantes:
                        writer.writerow({
                            'cedula': p['cedula'],
                            'nombre': p['nombre'],
                            'apellido': p['apellido'],
                            'direccion': p['direccion'],
                            'contacto': p['contacto'],
                            'estado_pago': p['estado_pago'],
                            'numero_dorsal': p['numero_dorsal']
                        })
            
            return True, f"✅ {len(participantes)} participantes exportados a {archivo_salida}"
        
        except Exception as e:
            return False, f"❌ Error al exportar: {str(e)}"
    
    def contar_participantes(self):
        """Retorna el número total de participantes registrados"""
        return len(self.obtener_todos())
    
    def estadisticas(self):
        """
        Genera estadísticas sobre los participantes
        
        Returns:
            dict: Estadísticas (total, pagados, pendientes, con_dorsal, sin_dorsal)
        """
        participantes = self.obtener_todos()
        total = len(participantes)
        pagados = sum(1 for p in participantes if p['estado_pago'] == 'PAGADO')
        pendientes = total - pagados
        con_dorsal = sum(1 for p in participantes if p['numero_dorsal'])
        sin_dorsal = total - con_dorsal
        
        return {
            'total': total,
            'pagados': pagados,
            'pendientes': pendientes,
            'con_dorsal': con_dorsal,
            'sin_dorsal': sin_dorsal
        }


# Ejemplo de uso
if __name__ == "__main__":
    gestor = GestorParticipantes()
    
    # Agregar participantes de ejemplo
    print("\n=== Agregando participantes de ejemplo ===")
    gestor.agregar_participante("12345678", "Juan", "Pérez", "Calle 123", "555-1234", "PAGADO", "001")
    gestor.agregar_participante("87654321", "María", "González", "Av. Principal 456", "555-5678", "PENDIENTE", "002")
    gestor.agregar_participante("11223344", "Carlos", "Ramírez", "Jr. Los Olivos 789", "555-9012", "PAGADO", "003")
    
    # Buscar por cédula
    print("\n=== Búsqueda por cédula ===")
    resultado = gestor.buscar_por_cedula("12345678")
    if resultado:
        print(f"Encontrado: {resultado['nombre']} {resultado['apellido']} - Dorsal: {resultado['numero_dorsal']}")
    
    # Mostrar estadísticas
    print("\n=== Estadísticas ===")
    stats = gestor.estadisticas()
    print(f"Total: {stats['total']}")
    print(f"Pagados: {stats['pagados']}")
    print(f"Pendientes: {stats['pendientes']}")
    print(f"Con dorsal asignado: {stats['con_dorsal']}")
    print(f"Sin dorsal: {stats['sin_dorsal']}")
